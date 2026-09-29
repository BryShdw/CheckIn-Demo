import io
from datetime import datetime
from flask import Blueprint, render_template, request, jsonify, send_file, redirect, url_for, flash
from flask_login import login_required, current_user
import openpyxl

from app.extensions import db, csrf
from app.models.user import User
from app.models.event import Event
from app.models.guest import Guest
from app.models.checkin import Checkin
from app.models.audit_log import AuditLog
from app.services.guest_service import (
    get_or_create_active_event,
    get_kiosk_settings,
    update_kiosk_settings,
    list_guests,
    create_guest,
    delete_guest,
    reset_all_checkins,
    get_event_stats,
    list_events,
    create_event,
    set_active_event,
    delete_event,
)
from app.services.printer_service import (
    get_printer_connection_status,
    purge_printer_queue,
    print_guest_ticket,
    build_label_image,
    DEFAULT_LABEL_TEMPLATE,
)
from app.services.face_service import delete_guest_face, delete_all_enrolled_faces
from app.services.import_service import sync_onedrive_link, parse_excel_file, parse_csv_stream, save_parsed_guests_to_event
from app.services.audit_service import log_audit

admin_bp = Blueprint("admin", __name__, url_prefix="/admin")
csrf.exempt(admin_bp)


@admin_bp.before_request
def check_auth_and_password():
    """Verifica autenticación y obliga al cambio de contraseña si must_change_password=True."""
    if not current_user.is_authenticated:
        return redirect(url_for("auth.login", next=request.path))

    if current_user.must_change_password and request.endpoint != "auth.change_password_view" and request.endpoint != "auth.logout":
        flash("Debes actualizar tu contraseña por defecto antes de continuar.", "warn")
        return redirect(url_for("auth.change_password_view"))


@admin_bp.route("/")
def admin_dashboard():
    event = get_or_create_active_event()
    return render_template("admin/admin.html", current_user=current_user, active_event=event)


# ── APIs Administrativas de Eventos ──────────────────────────────────────────

@admin_bp.route("/api/events", methods=["GET"])
def api_get_events():
    events = list_events()
    return jsonify({"status": "ok", "events": events})


@admin_bp.route("/api/events", methods=["POST"])
def api_create_event():
    data = request.get_json() or {}
    name = data.get("name", "").strip()
    if not name:
        return jsonify({"error": "El nombre del evento es requerido."}), 400

    location = data.get("location", "").strip()
    description = data.get("description", "").strip()
    set_active = bool(data.get("set_active", False))
    onedrive_url = data.get("onedrive_url", "").strip()

    event = create_event(
        name=name,
        location=location,
        description=description,
        set_active=set_active,
        onedrive_url=onedrive_url,
    )

    log_audit("EVENT_CREATE", user_id=current_user.id, target_type="EVENT", target_id=str(event.id),
              details=f"Creado evento '{event.name}'", ip_address=request.remote_addr)

    return jsonify({"status": "ok", "message": f"Evento '{event.name}' creado exitosamente.", "event": event.to_dict()}), 201


@admin_bp.route("/api/events/<int:event_id>/activate", methods=["POST"])
def api_activate_event(event_id):
    target = set_active_event(event_id)
    if not target:
        return jsonify({"error": "Evento no encontrado."}), 404

    log_audit("EVENT_ACTIVATE", user_id=current_user.id, target_type="EVENT", target_id=str(target.id),
              details=f"Activado evento '{target.name}'", ip_address=request.remote_addr)

    return jsonify({"status": "ok", "message": f"Evento '{target.name}' activado para acreditación.", "event": target.to_dict()})


@admin_bp.route("/api/events/<int:event_id>", methods=["DELETE"])
def api_delete_event(event_id):
    if not current_user.is_admin_or_super:
        return jsonify({"error": "Solo administradores pueden eliminar eventos."}), 403

    ok, msg = delete_event(event_id)
    if not ok:
        return jsonify({"error": msg}), 400

    log_audit("EVENT_DELETE", user_id=current_user.id, target_type="EVENT", target_id=str(event_id),
              details=msg, ip_address=request.remote_addr)

    return jsonify({"status": "ok", "message": msg})


@admin_bp.route("/api/import/onedrive", methods=["POST"])
def api_import_onedrive():
    data = request.get_json() or {}
    url = data.get("url", "").strip()
    event_id = data.get("event_id")

    event = Event.query.get(event_id) if event_id else get_or_create_active_event()
    if not event:
        return jsonify({"error": "Evento no encontrado."}), 404

    if not url:
        url = event.onedrive_url
    if not url:
        return jsonify({"error": "Ingresa un enlace de OneDrive o SharePoint válido."}), 400

    ok, msg, count = sync_onedrive_link(url, event.id)
    if not ok:
        return jsonify({"error": msg}), 400

    log_audit("IMPORT_ONEDRIVE", user_id=current_user.id, target_type="EVENT", target_id=str(event.id),
              details=f"Sincronizados {count} invitados desde OneDrive en evento '{event.name}'", ip_address=request.remote_addr)

    return jsonify({"status": "ok", "message": msg, "count": count, "event_id": event.id})


# ── APIs Administrativas de Usuarios ─────────────────────────────────────────

@admin_bp.route("/api/users", methods=["GET"])
def api_get_users():
    users = User.query.order_by(User.id.asc()).all()
    return jsonify({"status": "ok", "users": [u.to_dict() for u in users]})


@admin_bp.route("/api/users", methods=["POST"])
def api_create_user():
    if not current_user.is_admin_or_super:
        return jsonify({"error": "No tienes permisos para crear usuarios."}), 403

    data = request.get_json() or {}
    username = data.get("username", "").strip()
    password = data.get("password", "").strip()
    role = data.get("role", "OPERATOR").upper()
    email = data.get("email", "").strip() or None

    if not username or not password:
        return jsonify({"error": "Usuario y contraseña requeridos."}), 400

    if User.query.filter_by(username=username).first():
        return jsonify({"error": f"El nombre de usuario '{username}' ya está en uso."}), 409

    user = User(
        username=username,
        email=email,
        role=role if role in ("SUPERADMIN", "ADMIN", "OPERATOR") else "OPERATOR",
        is_active=True,
        must_change_password=False,
    )
    user.set_password(password)
    db.session.add(user)
    db.session.commit()

    log_audit("USER_CREATE", user_id=current_user.id, target_type="USER", target_id=str(user.id),
              details=f"Creado usuario '{username}' con rol '{role}'", ip_address=request.remote_addr)

    return jsonify({"status": "ok", "message": f"Usuario '{username}' creado exitosamente.", "user": user.to_dict()})


@admin_bp.route("/api/users/<int:user_id>", methods=["DELETE"])
def api_delete_user(user_id):
    if not current_user.is_superadmin:
        return jsonify({"error": "Solo un SuperAdmin puede eliminar usuarios."}), 403

    if user_id == current_user.id:
        return jsonify({"error": "No puedes eliminar tu propia cuenta de usuario."}), 400

    user = User.query.get_or_404(user_id)
    uname = user.username
    db.session.delete(user)
    db.session.commit()

    log_audit("USER_DELETE", user_id=current_user.id, target_type="USER", target_id=str(user_id),
              details=f"Eliminado usuario '{uname}'", ip_address=request.remote_addr)

    return jsonify({"status": "ok", "message": f"Usuario '{uname}' eliminado."})


# ── APIs Administrativas de Auditoría y Analítica ────────────────────────────

@admin_bp.route("/api/audit-logs", methods=["GET"])
def api_get_audit_logs():
    logs = AuditLog.query.order_by(AuditLog.id.desc()).limit(150).all()
    return jsonify({"status": "ok", "logs": [l.to_dict() for l in logs]})


@admin_bp.route("/api/analytics", methods=["GET"])
def api_get_analytics():
    event = get_or_create_active_event()
    stats = get_event_stats(event.id)

    # Agrupar checkins por intervalos de 1 hora
    checkins = Checkin.query.filter_by(event_id=event.id).order_by(Checkin.checked_in_at.asc()).all()
    timeline = {}
    for c in checkins:
        hour_key = c.checked_in_at.strftime("%H:00")
        timeline[hour_key] = timeline.get(hour_key, 0) + 1

    return jsonify({
        "status": "ok",
        "stats": stats,
        "timeline": {
            "labels": list(timeline.keys()),
            "data": list(timeline.values()),
        }
    })


# ── Importador de Archivos Excel (.xlsx / CSV) ───────────────────────────────

@admin_bp.route("/api/import/file", methods=["POST"])
def api_import_file():
    if "file" not in request.files:
        return jsonify({"error": "No se envió ningún archivo."}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Archivo no seleccionado."}), 400

    event = get_or_create_active_event()
    fname = file.filename.lower()
    raw_bytes = file.read()

    try:
        if fname.endswith(".xlsx"):
            parsed = parse_excel_file(raw_bytes)
        elif fname.endswith(".csv"):
            text = ""
            for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
                try:
                    text = raw_bytes.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            parsed = parse_csv_stream(text)
        else:
            return jsonify({"error": "Formato de archivo no soportado. Debe ser .xlsx o .csv."}), 400

        if not parsed:
            return jsonify({"error": "No se encontraron filas válidas o columnas reconocibles en el archivo."}), 400

        inserted, updated = save_parsed_guests_to_event(parsed, event)
        log_audit("IMPORT_FILE", user_id=current_user.id, target_type="EVENT", target_id=str(event.id),
                  details=f"Importado '{file.filename}': {inserted} nuevos, {updated} actualizados.", ip_address=request.remote_addr)

        return jsonify({
            "status": "ok",
            "message": f"Archivo '{file.filename}' procesado exitosamente ({len(parsed)} invitados: {inserted} nuevos, {updated} actualizados).",
            "total": len(parsed),
            "inserted": inserted,
            "updated": updated,
        })
    except Exception as e:
        return jsonify({"error": f"Error al procesar archivo: {e}"}), 500


# ── Exportar Asistencia a Excel (.xlsx) ──────────────────────────────────────

@admin_bp.route("/export-attendance")
def export_attendance():
    event = get_or_create_active_event()
    guests = Guest.query.filter_by(event_id=event.id, is_active=True).order_by(Guest.id.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Control de Asistencia"

    # Encabezados
    headers = ["Código", "Nombres y Apellidos", "Empresa", "Cargo", "Estado", "Hora Asistencia", "Método", "Categoría"]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")

    for g in guests:
        latest = g.latest_checkin
        ws.append([
            g.guest_code,
            g.full_name,
            g.company or "",
            g.position or "",
            "Registrado" if latest else "Pendiente",
            latest.checked_in_at.strftime("%d/%m/%Y %H:%M:%S") if latest else "—",
            latest.method if latest else "—",
            g.category,
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"Asistencia_{event.slug}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


@admin_bp.route("/api/reset-all", methods=["POST"])
@admin_bp.route("/api/reset-checkins", methods=["POST"])
def admin_api_reset_checkins():
    event = get_or_create_active_event()
    count = reset_all_checkins(event.id)
    user_id = current_user.id if current_user.is_authenticated else None
    log_audit("CHECKINS_RESET", user_id=user_id, target_type="EVENT", target_id=str(event.id),
              details=f"Reiniciadas {count} asistencias.", ip_address=request.remote_addr)
    return jsonify({
        "status": "ok",
        "success": True,
        "message": f"Se han reiniciado {count} registros de asistencia.",
        "reset_count": count,
    })


# ── Editor y Personalización de Plantillas de Credencial ────────────────────

def _get_sample_guest_data(event_id: int) -> dict:
    guest = Guest.query.filter_by(event_id=event_id, is_active=True).first()
    if guest:
        return {
            "id": guest.id,
            "guest_code": guest.guest_code,
            "full_name": guest.full_name,
            "name": guest.full_name,
            "company": guest.company or "EMPRESA DE EJEMPLO",
            "position": guest.position or "DIRECTOR EJECUTIVO",
            "category": guest.category or "VIP",
            "event_id": event_id,
        }
    return {
        "id": 1,
        "guest_code": "ASIS-0001",
        "full_name": "CARLOS EDUARDO MENDOZA SILVA",
        "name": "CARLOS EDUARDO MENDOZA SILVA",
        "company": "CORPORACIÓN TECNOLÓGICA",
        "position": "GERENTE DE OPERACIONES",
        "category": "VIP",
        "event_id": event_id,
    }


@admin_bp.route("/api/label-template", methods=["GET"])
def api_get_label_template():
    event = get_or_create_active_event()
    template = event.get_label_template()
    sample_guest = _get_sample_guest_data(event.id)
    return jsonify({
        "status": "ok",
        "template": template,
        "default_template": DEFAULT_LABEL_TEMPLATE,
        "sample_guest": sample_guest,
    })


@admin_bp.route("/api/label-template", methods=["POST"])
def api_save_label_template():
    event = get_or_create_active_event()
    data = request.get_json() or {}
    template = data.get("template")
    if not isinstance(template, dict):
        return jsonify({"error": "Estructura de plantilla no válida."}), 400

    event.set_label_template(template)
    db.session.commit()

    user_id = current_user.id if current_user.is_authenticated else None
    log_audit("LABEL_TEMPLATE_UPDATE", user_id=user_id, target_type="EVENT", target_id=str(event.id),
              details=f"Plantilla de credencial actualizada para evento '{event.name}'", ip_address=request.remote_addr)

    return jsonify({
        "status": "ok",
        "message": "Plantilla de credencial guardada correctamente.",
        "template": event.get_label_template(),
    })


@admin_bp.route("/api/label-preview", methods=["POST"])
def api_preview_label():
    event = get_or_create_active_event()
    data = request.get_json() or {}
    template = data.get("template")
    guest_data = data.get("guest") or _get_sample_guest_data(event.id)

    img = build_label_image(guest_data, template=template)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return send_file(buf, mimetype="image/png")


@admin_bp.route("/api/label-template/test-print", methods=["POST"])
def api_test_print_label():
    event = get_or_create_active_event()
    data = request.get_json() or {}
    template = data.get("template")
    guest_data = data.get("guest") or _get_sample_guest_data(event.id)

    print_ok, print_msg = print_guest_ticket(guest_data, template=template)
    user_id = current_user.id if current_user.is_authenticated else None
    log_audit("LABEL_TEST_PRINT", user_id=user_id, target_type="EVENT", target_id=str(event.id),
              details=f"Prueba de impresión de credencial. Resultado: {print_msg}", ip_address=request.remote_addr)

    return jsonify({
        "status": "ok" if print_ok else "error",
        "success": print_ok,
        "message": print_msg,
    })

